"""CE-S5: build the Delegation of Authority fixture variants from the baseline Master Control Workbook.

The baseline (NEX_BK-SYS-003_Master_Control_Workbook.xlsx) is synthetic. Each variant changes one thing so
the CE-S5 smoke can prove one behaviour:

  NEX_BK-SYS-003_approval_limits_changed.xlsx    Purchases Tier 1 limit 5000 -> 8000        (re-import = new version)
  NEX_BK-SYS-003_approval_limits_bad_rows.xlsx   four unusable rows added                    (row-level rejections)
  NEX_BK-SYS-003_no_approval_limits_sheet.xlsx   the Approval Limits sheet removed           (whole-file rejection)
  NEX_BK-SYS-003_header_changed.xlsx             a column heading renamed                    (STRUCTURE_CHANGED)

Usage: python build_delegation_variants.py [--dir <folder holding the baseline>]
"""
import argparse
from pathlib import Path

import openpyxl

BASE = "NEX_BK-SYS-003_Master_Control_Workbook.xlsx"


def load(folder):
    return openpyxl.load_workbook(Path(folder) / BASE)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dir", default=str(Path(__file__).resolve().parent))
    folder = Path(ap.parse_args().dir)

    wb = load(folder)
    wb["Approval Limits"].cell(4, 2, 8000)
    wb.save(folder / "NEX_BK-SYS-003_approval_limits_changed.xlsx")

    wb = load(folder)
    ws = wb["Approval Limits"]
    ws.append(["Loans", "abc", "Finance Manager", 10000, "Authorised Manager", "Owner / Board"])            # unreadable limit
    ws.append(["Refunds", 5000, "Finance Manager", 2000, "Authorised Manager", "Owner / Board"])            # Tier 2 below Tier 1
    ws.append(["Purchases", 1, "Finance Manager", 2, "Authorised Manager", "Owner / Board"])                # duplicate type
    ws.append(["Petty cash", 0.05, "Finance Manager", 0.1, "Authorised Manager", "Owner / Board"])          # ambiguous (percent cell?)
    wb.save(folder / "NEX_BK-SYS-003_approval_limits_bad_rows.xlsx")

    wb = load(folder)
    del wb["Approval Limits"]
    wb.save(folder / "NEX_BK-SYS-003_no_approval_limits_sheet.xlsx")

    wb = load(folder)
    wb["Approval Limits"].cell(3, 2, "First Limit")
    wb.save(folder / "NEX_BK-SYS-003_header_changed.xlsx")
    print("built 4 delegation variants in", folder)


if __name__ == "__main__":
    main()
